# -*- coding: utf-8 -*-
"""NEIS '과목별세부능력및특기사항' 익명화 예시 양식 생성기.

실제 NEIS 다운로드 파일의 **헤더 구조만** 재현하고, 학생 데이터(학생개인번호·과목코드·
반/번호·성명)는 전부 **더미 값**으로 채운다. 개인정보가 전혀 없으므로 스킬 assets에
안전하게 둘 수 있다. build_gwagi.py의 NEIS 탐지/채우기 로직의 회귀 픽스처로 쓴다.

세특 칸(J열)은 비워 둔다 — 채우기 모드가 여기에 평가문을 넣는다.
재생성: python make_neis_sample.py
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "NEIS_과목세특_양식_예시.xlsx")

# 실제 NEIS 다운로드 파일의 1행 헤더(순서·문자열 그대로). 탐지 로직의 기준이 된다.
HEADERS = [
    "학년도", "학기", "학년", "학생개인번호", "과목", "과목코드",
    "반/번호", "성명", "학적변동\r\n구분", "세부능력 및 특기사항",
    "영재·발명교육 기록사항",
]
SETUK_COL = 10  # J = 세부능력 및 특기사항 (채우기 모드가 채울 칸 → 예시에선 비움)

# 더미 학생 5명. 학생개인번호·성명은 가공값(실제 인물 아님).
DUMMY = [
    # 학년도 학기 학년 학생개인번호 과목 과목코드   반/번호 성명     학적변동
    ("2026", "1", "1", "0000000001", "정보", "0000000000", "1/1",  "홍길동", "재학"),
    ("2026", "1", "1", "0000000002", "정보", "0000000000", "1/2",  "김철수", "재학"),
    ("2026", "1", "1", "0000000003", "정보", "0000000000", "1/3",  "이영희", "재학"),
    ("2026", "1", "1", "0000000004", "정보", "0000000000", "1/4",  "박민수", "재학"),
    ("2026", "1", "1", "0000000005", "정보", "0000000000", "1/5",  "최지은", "재학"),
]

THIN = Border(*[Side(style="thin", color="C0C0C0")] * 4)
HEAD_FILL = PatternFill("solid", fgColor="E2EFDA")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "과목별세부능력및특기사항"

    for c, h in enumerate(HEADERS, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
        cell.alignment = CENTER
        cell.border = THIN

    for r, row in enumerate(DUMMY, 2):
        for c in range(1, len(HEADERS) + 1):
            val = row[c - 1] if c - 1 < len(row) else ""
            # 세특(J)·영재(K)는 비워 둔다
            if c >= SETUK_COL:
                val = ""
            cell = ws.cell(r, c, val)
            cell.border = THIN
            cell.alignment = CENTER if c != SETUK_COL else WRAP

    widths = {4: 14, 6: 12, 7: 9, 8: 10, 9: 10, 10: 60, 11: 18}
    for c in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(c, 8)
    ws.freeze_panes = "A2"

    wb.save(OUT)
    print(f"저장: {OUT}  (헤더 {len(HEADERS)}열, 더미 {len(DUMMY)}명, 세특칸 J는 빈칸)")


if __name__ == "__main__":
    main()
