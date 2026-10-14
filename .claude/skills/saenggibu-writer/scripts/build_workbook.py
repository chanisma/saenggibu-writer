#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
students.json  ->  통합 Excel(.xlsx)

한 시트에 학생별 자료를 활동별 열로 정리하고, 평가문 / 바이트수 / 목표·판정을 함께 담는다.
바이트수가 '목표바이트' 범위를 벗어나면 셀을 색으로 경고한다.

students.json 구조(예):
{
  "subject": "인공지능 기초",
  "students": [
    {
      "번호": "10101", "이름": "홍길동", "성적": 92,
      "제약사항": "결석 잦음",
      "활동": {
        "이미지분류프로젝트": "....",
        "AI서비스개발": "....",
        "자기평가": "...."
      },
      "목표바이트": "1400~1500",
      "계획": {"주장": "...", "근거": [...], "순서": "...", "제외": "..."},   // 4단계에서 채움
      "평가문": "....",      // 5단계에서 채움
      "바이트수": 1452        // 5단계에서 채움
    }
  ]
}

사용법: python build_workbook.py students.json 결과.xlsx
"""
import sys
import os
import json

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from byte_count import count_bytes, parse_target  # noqa: E402
from evaltext import base_text, review_text, has_review, has_plan, plan_text  # noqa: E402
import rawdata  # noqa: E402

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl 가 필요합니다:  pip install openpyxl")

HEADER_FILL = PatternFill("solid", fgColor="2F5496")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
OVER_FILL = PatternFill("solid", fgColor="F4CCCC")   # 초과: 붉은 톤
UNDER_FILL = PatternFill("solid", fgColor="FFF2CC")  # 미달: 노란 톤
OK_FILL = PatternFill("solid", fgColor="D9EAD3")     # 적정: 녹색 톤
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Border(*[Side(style="thin", color="D0D0D0")] * 4)


def collect_activity_keys(students):
    keys = []
    for s in students:
        for k in (s.get("활동") or {}):
            if k not in keys:
                keys.append(k)
    return keys


def main():
    if len(sys.argv) < 3:
        sys.exit("사용법: python build_workbook.py students.json 결과.xlsx")
    src, out = sys.argv[1], sys.argv[2]
    with open(src, encoding="utf-8") as f:
        data = json.load(f)
    students = data.get("students", data if isinstance(data, list) else [])
    # 활동 원문은 students.json에 복사해 두지 않고 `원자료` 파일에서 가져온다(3단계 산출·불변)
    nraw, rawp = rawdata.attach(src, data)
    if nraw:
        print(f"원자료 연결: {os.path.basename(rawp)} → 학생 {nraw}명의 활동 원문 사용")
    subject = data.get("subject", "")

    act_keys = collect_activity_keys(students)
    base_cols = ["번호", "이름", "성적", "제약사항"]
    # 구성 계획(4단계)이 있으면 자료와 평가문 사이에 둔다 — 근거에서 계획을 거쳐 글로 읽히는 순서
    hasplan = has_plan(students)
    tail_cols = (["구성계획"] if hasplan else []) + ["목표바이트", "평가문", "바이트수", "판정"]
    # 품질 검토 패스를 돌렸으면 원본을 덮지 않고 검토본 열을 뒤에 덧붙인다.
    hasrev = has_review(students)
    if hasrev:
        tail_cols += ["평가문_검토본", "바이트수_검토본", "판정_검토본"]
    headers = base_cols + act_keys + tail_cols

    wb = Workbook()
    ws = wb.active
    ws.title = (subject or "생활기록부")[:30]

    # 헤더
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN

    # 본문
    for r, s in enumerate(students, 2):
        # 저장된 '바이트수'는 쉼표 정리 등으로 평가문이 바뀐 뒤 낡았을 수 있어 항상 재계산한다.
        lo, hi = parse_target(s.get("목표바이트", ""))

        def judge(t):
            n = count_bytes(t)
            if not t:
                return n, "", None
            if hi is not None and n > hi:
                return n, f"초과(+{n - hi})", OVER_FILL
            if lo is not None and n < lo:
                return n, f"미달(-{lo - n})", UNDER_FILL
            return n, "적정", OK_FILL

        text = base_text(s)
        nbytes, verdict, fill = judge(text)
        rtext = review_text(s)
        rbytes, rverdict, rfill = judge(rtext)

        row = {
            "번호": s.get("번호", ""),
            "이름": s.get("이름", ""),
            "성적": s.get("성적", ""),
            "제약사항": s.get("제약사항", ""),
            "구성계획": plan_text(s),
            "목표바이트": s.get("목표바이트", ""),
            "평가문": text,
            "바이트수": nbytes if text else "",
            "판정": verdict,
            "평가문_검토본": rtext,
            "바이트수_검토본": rbytes if rtext else "",
            "판정_검토본": rverdict,
        }
        for k in act_keys:
            row[k] = (s.get("활동") or {}).get(k, "")

        for c, h in enumerate(headers, 1):
            val = row.get(h, "")
            cell = ws.cell(r, c, val)
            if isinstance(val, str) and val.startswith("="):
                cell.data_type = "s"  # '=' 시작 학생 자료가 수식으로 저장되지 않게
            cell.border = THIN
            if h in ("평가문", "평가문_검토본", "구성계획") or h in act_keys or h == "제약사항":
                cell.alignment = WRAP
            elif h in ("번호", "이름", "성적", "바이트수", "목표바이트", "판정",
                       "바이트수_검토본", "판정_검토본"):
                cell.alignment = CENTER
            if h in ("바이트수", "판정") and fill is not None:
                cell.fill = fill
            if h in ("바이트수_검토본", "판정_검토본") and rfill is not None:
                cell.fill = rfill

    # 열 너비
    widths = {"번호": 9, "이름": 9, "성적": 7, "제약사항": 18, "구성계획": 46,
              "목표바이트": 12, "평가문": 70, "바이트수": 9, "판정": 11,
              "평가문_검토본": 70, "바이트수_검토본": 9, "판정_검토본": 11}
    for c, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(h, 32)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    wb.save(out)
    print(f"저장 완료: {out}  (학생 {len(students)}명, 활동열 {len(act_keys)}개)")


if __name__ == "__main__":
    main()
